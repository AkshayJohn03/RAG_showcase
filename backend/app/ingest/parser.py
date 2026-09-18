"""Document parser: .md/.txt natively, .pdf/.docx if libs present."""
from pathlib import Path

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    import docx
except Exception:
    docx = None


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt", ".markdown"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf" and PdfReader is not None:
        reader = PdfReader(str(path))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if suffix == ".docx" and docx is not None:
        doc = docx.Document(str(path))
        out = [p.text for p in doc.paragraphs]
        for t in doc.tables:  # multimodal: keep tables as markdown
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
                for r in rows[1:]:
                    out.append("| " + " | ".join(r) + " |")
        return "\n".join(out)
    # graceful: unknown type or missing dep
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def load_raw(raw_dir: Path) -> list[dict]:
    docs = []
    for p in sorted(raw_dir.rglob("*")):
        if p.is_file() and not p.name.startswith("."):
            text = parse_file(p)
            if text.strip():
                docs.append({"doc_id": p.relative_to(raw_dir).as_posix(), "source": p.name, "text": text})
    return docs
